"""
zebra_excel_annex.py

Genera el Excel anexo de Proyección de Inversión para acompañar la propuesta DOCX.
Replica la estructura de las dos hojas del DEMO original:
  - Hoja 1: Calculadora de Inversión (con inputs editables y fórmulas)
  - Hoja 2: Plan de Inversión, Proyección a 12 meses

Estilo: alineado al Zebra Design System (ink monocromático, sin amarillo de
chrome, hairline borders, eyebrows en JetBrains Mono, KPI principal con fondo
ink).
"""

from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter


# =============================================================================
# Paleta Zebra (Design System v2.0, tema light)
# =============================================================================

INK         = "0A0A0A"  # --text / --accent
INK_700     = "262626"  # hover, secondary emphasis
INK_500     = "737373"  # muted text (placeholders, hints, AA legible)
INK_300     = "CBD5E1"  # light borders
INK_200     = "E5E5E5"  # default borders (hairline)
INK_100     = "F5F7FA"  # subtle backgrounds
INK_50      = "F9FAFB"  # --surface-2 (recessed sections)
BONE        = "FFFFFF"  # --surface

# Fuentes del DS. Si el sistema no las tiene, Excel/Sheets caen a Calibri /
# Roboto, ambas aceptables. Google Sheets sí tiene Inter por default.
SANS_FAMILY = "Inter"
MONO_FAMILY = "JetBrains Mono"


# =============================================================================
# Helpers de estilo
# =============================================================================

def _hairline(top=False, bottom=False, left=False, right=False, color=INK_200):
    side = Side(border_style="thin", color=color)
    return Border(
        top=side if top else None,
        bottom=side if bottom else None,
        left=side if left else None,
        right=side if right else None,
    )


def _full_hairline(color=INK_200):
    side = Side(border_style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def _fill(color):
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _label_cell(cell, text, bold=False, size=10, color=INK, fill_color=None,
                family=SANS_FAMILY, align="left"):
    cell.value = text
    cell.font = Font(name=family, size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    if fill_color:
        cell.fill = _fill(fill_color)


def _value_cell(cell, value, bold=False, size=10, color=INK,
                fill_color=None, fmt=None, align="right", family=MONO_FAMILY,
                border=None):
    cell.value = value
    cell.font = Font(name=family, size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    if fill_color:
        cell.fill = _fill(fill_color)
    if fmt:
        cell.number_format = fmt
    if border is not None:
        cell.border = border


def _eyebrow(cell, text):
    """Section label: JetBrains Mono, 9pt, uppercase, tracking implícito, muted."""
    cell.value = text.upper()
    cell.font = Font(name=MONO_FAMILY, size=9, bold=True, color=INK_500)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def _title_cell(cell, text, size=14):
    """Title: Inter semibold, ink color, sin fondo."""
    cell.value = text
    cell.font = Font(name=SANS_FAMILY, size=size, bold=True, color=INK)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def _subtitle_cell(cell, text):
    cell.value = text
    cell.font = Font(name=SANS_FAMILY, size=9, color=INK_500, italic=False)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def _input_cell(cell, value, fmt=None):
    """Celda editable (input del modelo). Fondo recesado + hairline ink."""
    cell.value = value
    cell.font = Font(name=MONO_FAMILY, size=10, bold=True, color=INK)
    cell.alignment = Alignment(horizontal="right", vertical="center")
    cell.fill = _fill(INK_50)
    cell.border = _full_hairline(INK_300)
    if fmt:
        cell.number_format = fmt


def _output_cell(cell, value, fmt=None, bold=True):
    """Celda calculada. Fondo bone, hairline ink-200."""
    cell.value = value
    cell.font = Font(name=MONO_FAMILY, size=10, bold=bold, color=INK)
    cell.alignment = Alignment(horizontal="right", vertical="center")
    cell.fill = _fill(BONE)
    cell.border = _full_hairline(INK_200)
    if fmt:
        cell.number_format = fmt


def _kpi_cell(cell, value, fmt=None):
    """KPI principal. Fondo ink (#0A0A0A), texto blanco. Stat emphasis card."""
    cell.value = value
    cell.font = Font(name=MONO_FAMILY, size=11, bold=True, color=BONE)
    cell.alignment = Alignment(horizontal="right", vertical="center")
    cell.fill = _fill(INK)
    if fmt:
        cell.number_format = fmt


def _kpi_label(cell, text):
    cell.value = text
    cell.font = Font(name=SANS_FAMILY, size=11, bold=True, color=BONE)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.fill = _fill(INK)


def _table_header(cell, text):
    """Header de tabla: fondo ink, texto blanco, mono, centrado."""
    cell.value = text
    cell.font = Font(name=MONO_FAMILY, size=9, bold=True, color=BONE)
    cell.fill = _fill(INK)
    cell.alignment = Alignment(horizontal="center", vertical="center")


# =============================================================================
# Hoja 1, Calculadora de Inversión
# =============================================================================

def _build_calc_sheet(ws, calc_output):
    from zebra_investment_calc import (
        LEADS_DIA_POR_ASESOR, DIAS_PROMEDIO_MES,
    )

    ws.title = "Calculadora"
    ws.sheet_view.showGridLines = False

    # Anchos
    widths = {"A": 2, "B": 36, "C": 4, "D": 22, "E": 2, "F": 30, "G": 16}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # Título y subtítulo
    _title_cell(ws["B2"], "Calculadora de Inversión", size=18)
    _subtitle_cell(
        ws["B3"],
        "Ingeniería inversa: valor de proyecto, absorción y pauta requerida"
    )

    # ===== INPUTS =====
    _eyebrow(ws["B5"], "Inputs del modelo")

    input_rows = [
        ("Unidades Totales", calc_output.unidades_totales, "0", 6),
        ("Ticket Promedio", calc_output.ticket_promedio, '"$"#,##0" MXN"', 7),
        ("Absorción del Proyecto (meses)", calc_output.absorcion_meses, "0", 8),
        ("CPL Estimado", calc_output.cpl, '"$"#,##0" MXN"', 9),
        ("% Inversión sobre Valor", calc_output.porcentaje_inversion, "0.00%", 10),
    ]
    for label, value, fmt, row in input_rows:
        _label_cell(ws[f"B{row}"], label, size=10)
        _input_cell(ws[f"D{row}"], value, fmt=fmt)

    # ===== TASAS =====
    _eyebrow(ws["B12"], "Tasas de conversión")

    tasa_rows = [
        ("% Conversión a Cita", calc_output.tasas['cita'], 13),
        ("% Asistencia", calc_output.tasas['asistencia'], 14),
        ("% Apartado", calc_output.tasas['apartado'], 15),
        ("% Cierre", calc_output.tasas['cierre'], 16),
    ]
    for label, value, row in tasa_rows:
        _label_cell(ws[f"B{row}"], label, size=10)
        _input_cell(ws[f"D{row}"], value, fmt="0.0%")

    # ===== OUTPUTS CALCULADOS =====
    _eyebrow(ws["B18"], "Resultados calculados")

    output_rows = [
        ("Valor del Proyecto", "=D6*D7", '"$"#,##0" MXN"', 19),
        ("CPA Estimado", "=D10*D7", '"$"#,##0" MXN"', 20),
        ("Presupuesto Total", "=D19*D10", '"$"#,##0" MXN"', 21),
        ("ROA Esperado", "=D19/D21", '0.0"x"', 22),
        ("Unidades Objetivo Mensuales", "=D6/D8", "0.00", 23),
    ]
    for label, formula, fmt, row in output_rows:
        _label_cell(ws[f"B{row}"], label, size=10)
        _output_cell(ws[f"D{row}"], formula, fmt=fmt)

    # KPI principal (Inversión pauta) — ink fill, white text
    _kpi_label(ws["B24"], "Inversión Pauta Mensual")
    _kpi_cell(ws["D24"], "=D21/D8", fmt='"$"#,##0" MXN"')

    # Resto de outputs
    tail_rows = [
        ("Leads Esperados (mensuales)", "=ROUNDDOWN(D24/D9,0)", "#,##0", 25),
        ("Capacidad Máx. Mensual / Asesor",
         f"=ROUNDDOWN({LEADS_DIA_POR_ASESOR}*{DIAS_PROMEDIO_MES},0)", "0", 26),
        ("Sala Necesaria (asesores)", "=ROUNDDOWN(D25/D26,0)", "0", 27),
    ]
    for label, formula, fmt, row in tail_rows:
        _label_cell(ws[f"B{row}"], label, size=10, bold=(row == 27))
        _output_cell(ws[f"D{row}"], formula, fmt=fmt)

    # ===== CASCADA DEL FUNNEL (columna derecha) =====
    _eyebrow(ws["F5"], "Eficiencia comercial esperada")

    funnel_rows = [
        ("Citas (mes)", "=D25*D13", 7),
        ("Asistencias (mes)", "=G7*D14", 8),
        ("Apartados (mes)", "=G8*D15", 9),
        ("Cierres (mes)", "=G9*D16", 10),
    ]
    for label, formula, row in funnel_rows:
        _label_cell(ws[f"F{row}"], label, size=10)
        _output_cell(ws[f"G{row}"], formula, fmt="0", bold=(label == "Cierres (mes)"))

    # ===== NOTAS =====
    ws.merge_cells("B29:G31")
    _label_cell(
        ws["B29"],
        "Notas: el % de inversión sobre valor del proyecto es un promedio "
        "fijo (3%) como punto de partida; se ajusta según ticket, ubicación "
        "y dinámica del proyecto. CPL y tasas de conversión basadas en "
        "benchmarks Zebra Real Estate; pueden ajustarse al rendimiento "
        "real del cliente.",
        size=8, color=INK_500
    )
    ws["B29"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Altura de filas
    ws.row_dimensions[2].height = 28
    ws.row_dimensions[3].height = 16
    for r in list(range(5, 29)):
        ws.row_dimensions[r].height = 22
    ws.row_dimensions[24].height = 26  # KPI con un poco más de aire


# =============================================================================
# Hoja 2, Plan 12 meses
# =============================================================================

def _build_plan_sheet(ws, calc_output, plan_output):
    from zebra_investment_calc import (
        LEADS_DIA_POR_ASESOR, DIAS_PROMEDIO_MES,
        VENTANA_CITA_DIAS, VENTANA_APARTADO_DIAS, VENTANA_CIERRE_DIAS,
        MULTIPLICADORES_UNIDADES,
    )

    ws.title = "Plan 12 Meses"
    ws.sheet_view.showGridLines = False

    # Anchos
    ws.column_dimensions["A"].width = 36
    for c in range(2, 19):
        ws.column_dimensions[get_column_letter(c)].width = 12

    # Título
    _title_cell(ws["A1"], "Plan de Inversión: Proyección a 12 meses", size=16)
    _subtitle_cell(
        ws["A2"],
        "Zebra High Performance Marketing · Implementación + Operación · "
        "6 ventanas de conversión por mes"
    )

    # ===== SUPUESTOS =====
    _eyebrow(ws["A4"], "Supuestos del modelo")

    supuestos = [
        ("CPL Proyectado", calc_output.cpl, '"$"#,##0" MXN"'),
        ("Leads / día / asesor", LEADS_DIA_POR_ASESOR, "0"),
        ("Días promedio del mes", DIAS_PROMEDIO_MES, "0.0"),
        ("% Conversión a Cita", calc_output.tasas['cita'], "0.0%"),
        ("% Asistencia", calc_output.tasas['asistencia'], "0.0%"),
        ("% Apartado", calc_output.tasas['apartado'], "0.0%"),
        ("% Cierre", calc_output.tasas['cierre'], "0.0%"),
        ("Ventana Cita (días)", VENTANA_CITA_DIAS, "0"),
        ("Ventana Apartado (días)", VENTANA_APARTADO_DIAS, "0"),
        ("Ventana Cierre (días)", VENTANA_CIERRE_DIAS, "0"),
        ("Precio por propiedad base", calc_output.ticket_promedio, '"$"#,##0" MXN"'),
    ]
    row = 5
    for label, value, fmt in supuestos:
        _label_cell(ws.cell(row=row, column=1), label, size=9)
        _input_cell(ws.cell(row=row, column=2), value, fmt=fmt)
        row += 1

    # ===== PROYECCIÓN MENSUAL =====
    start_row = 18
    _eyebrow(ws.cell(row=start_row, column=1), "Proyección mensual: 12 meses / 4 trimestres")

    hr = start_row + 1
    headers = (
        ["CONCEPTO"]
        + [f"M{i}" for i in range(1, 4)] + ["Q1"]
        + [f"M{i}" for i in range(4, 7)] + ["Q2"]
        + [f"M{i}" for i in range(7, 10)] + ["Q3"]
        + [f"M{i}" for i in range(10, 13)] + ["Q4"]
        + ["TOTAL"]
    )
    for i, h in enumerate(headers, start=1):
        _table_header(ws.cell(row=hr, column=i), h)

    ws.row_dimensions[hr].height = 22

    # ===== EQUIPO E INVERSIÓN =====
    cur = hr + 2
    _eyebrow(ws.cell(row=cur, column=1), "Equipo e inversión")
    cur += 1

    # Asesores activos
    _label_cell(ws.cell(row=cur, column=1), "Asesores activos", size=9)
    for i, m in enumerate(plan_output.meses):
        col = _month_to_col(i + 1)
        _value_cell(ws.cell(row=cur, column=col), m.asesores_activos,
                    fmt="0", align="center")
    for q_col, q in zip([5, 9, 13, 17], plan_output.trimestres):
        _value_cell(ws.cell(row=cur, column=q_col), q.asesores_promedio,
                    bold=True, fmt="0.0", align="center", fill_color=INK_50)
    _value_cell(ws.cell(row=cur, column=18), plan_output.asesores_promedio_anual,
                bold=True, fmt="0.0", align="center", fill_color=INK_50)
    cur += 1

    # Leads
    _label_cell(ws.cell(row=cur, column=1), "Leads generados", size=9)
    for i, m in enumerate(plan_output.meses):
        col = _month_to_col(i + 1)
        _value_cell(ws.cell(row=cur, column=col), int(m.leads_generados),
                    fmt="#,##0", align="center")
    for q_col, q in zip([5, 9, 13, 17], plan_output.trimestres):
        _value_cell(ws.cell(row=cur, column=q_col), int(q.leads_total),
                    bold=True, fmt="#,##0", align="center", fill_color=INK_50)
    _value_cell(ws.cell(row=cur, column=18), int(plan_output.leads_total_anual),
                bold=True, fmt="#,##0", align="center", fill_color=INK_50)
    cur += 1

    # Inversión pauta: trimestres y total en ink (stat emphasis), meses en bone
    _label_cell(ws.cell(row=cur, column=1), "Inversión en pauta", size=9, bold=True)
    for i, m in enumerate(plan_output.meses):
        col = _month_to_col(i + 1)
        _value_cell(ws.cell(row=cur, column=col), m.inversion_pauta,
                    fmt='"$"#,##0', align="center")
    for q_col, q in zip([5, 9, 13, 17], plan_output.trimestres):
        _value_cell(ws.cell(row=cur, column=q_col), q.inversion_total,
                    bold=True, fmt='"$"#,##0', align="center",
                    fill_color=INK, color=BONE)
    _value_cell(ws.cell(row=cur, column=18), plan_output.inversion_total_anual,
                bold=True, fmt='"$"#,##0', align="center",
                fill_color=INK, color=BONE)
    cur += 2

    # ===== EMBUDO REAL =====
    _eyebrow(ws.cell(row=cur, column=1), "Embudo real (con ventanas de conversión)")
    cur += 1

    funnel_rows = [
        ("Citas reales", "citas_reales", "0"),
        ("Asistencias reales", "asistencias_reales", "0"),
        ("Apartados reales", "apartados_reales", "0"),
        ("Cierres reales", "cierres_reales", "0"),
    ]
    for label, attr, fmt in funnel_rows:
        is_cierres = (attr == "cierres_reales")
        _label_cell(ws.cell(row=cur, column=1), label, size=9, bold=is_cierres)
        for i, m in enumerate(plan_output.meses):
            col = _month_to_col(i + 1)
            val = getattr(m, attr)
            _value_cell(ws.cell(row=cur, column=col), val,
                        fmt=fmt, align="center")
        for q_col, q in zip([5, 9, 13, 17], plan_output.trimestres):
            val = getattr(q, attr)
            _value_cell(ws.cell(row=cur, column=q_col), val,
                        bold=True, fmt=fmt, align="center", fill_color=INK_50)
        total = sum(getattr(m, attr) for m in plan_output.meses)
        _value_cell(ws.cell(row=cur, column=18), total,
                    bold=True, fmt=fmt, align="center", fill_color=INK_50)
        cur += 1
    cur += 1

    # ===== INGRESOS POR ESCENARIO =====
    _eyebrow(ws.cell(row=cur, column=1),
             "Proyección de ingresos (cierres × ticket × multiplicador)")
    cur += 1

    for mult in MULTIPLICADORES_UNIDADES:
        _label_cell(ws.cell(row=cur, column=1),
                    f"Ingresos @ {mult:.1f} unidad/operación", size=9)
        for i, m in enumerate(plan_output.meses):
            col = _month_to_col(i + 1)
            _value_cell(ws.cell(row=cur, column=col), m.ingresos_por_escenario[mult],
                        fmt='"$"#,##0', align="center")
        for q_col, q in zip([5, 9, 13, 17], plan_output.trimestres):
            _value_cell(ws.cell(row=cur, column=q_col), q.ingresos_por_escenario[mult],
                        bold=True, fmt='"$"#,##0', align="center", fill_color=INK_50)
        _value_cell(ws.cell(row=cur, column=18),
                    plan_output.ingresos_anual_por_escenario[mult],
                    bold=True, fmt='"$"#,##0', align="center", fill_color=INK_50)
        cur += 1
    cur += 1

    # ===== MARGEN =====
    _eyebrow(ws.cell(row=cur, column=1),
             "Margen de adquisición (ingresos menos inversión)")
    cur += 1

    for mult in MULTIPLICADORES_UNIDADES:
        _label_cell(ws.cell(row=cur, column=1),
                    f"Margen @ {mult:.1f} unidad/operación", size=9)
        for i, m in enumerate(plan_output.meses):
            col = _month_to_col(i + 1)
            _value_cell(ws.cell(row=cur, column=col), m.margen_por_escenario[mult],
                        fmt='"$"#,##0', align="center")
        for q_col, q in zip([5, 9, 13, 17], plan_output.trimestres):
            _value_cell(ws.cell(row=cur, column=q_col), q.margen_por_escenario[mult],
                        bold=True, fmt='"$"#,##0', align="center", fill_color=INK_50)
        _value_cell(ws.cell(row=cur, column=18),
                    plan_output.margen_anual_por_escenario[mult],
                    bold=True, fmt='"$"#,##0', align="center", fill_color=INK_50)
        cur += 1
    cur += 1

    # ===== ROA =====
    _eyebrow(ws.cell(row=cur, column=1), "ROA acumulado (ingresos / inversión)")
    cur += 1

    for mult in MULTIPLICADORES_UNIDADES:
        _label_cell(ws.cell(row=cur, column=1),
                    f"ROA @ {mult:.1f} unidad/operación", size=9)
        for i, m in enumerate(plan_output.meses):
            col = _month_to_col(i + 1)
            _value_cell(ws.cell(row=cur, column=col), m.roa_por_escenario[mult],
                        fmt='0.0"x"', align="center")
        for q_col, q in zip([5, 9, 13, 17], plan_output.trimestres):
            _value_cell(ws.cell(row=cur, column=q_col), q.roa_por_escenario[mult],
                        bold=True, fmt='0.0"x"', align="center", fill_color=INK_50)
        _value_cell(ws.cell(row=cur, column=18),
                    plan_output.roa_anual_por_escenario[mult],
                    bold=True, fmt='0.0"x"', align="center", fill_color=INK_50)
        cur += 1


def _month_to_col(month):
    """Mapea mes (1-12) a columna del Excel:
    M1=2, M2=3, M3=4, Q1=5, M4=6, M5=7, M6=8, Q2=9, ...
    """
    if month <= 3:
        return month + 1
    elif month <= 6:
        return month + 2
    elif month <= 9:
        return month + 3
    else:
        return month + 4


def generate_investment_excel(calc_output, plan_output, output_path):
    """
    Genera el Excel anexo con dos hojas:
      - Calculadora de Inversión (con fórmulas vinculadas, inputs editables)
      - Plan de Inversión 12 meses (con valores calculados)
    """
    wb = Workbook()
    ws1 = wb.active
    _build_calc_sheet(ws1, calc_output)

    ws2 = wb.create_sheet("Plan 12 Meses")
    _build_plan_sheet(ws2, calc_output, plan_output)

    wb.save(output_path)
    return output_path
